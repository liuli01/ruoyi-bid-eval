"""评审意见 Word 导出

生成会商意见函 .docx，格式：
1. 项目信息头部
2. 风险分类分段（按 department 分组）
3. 各条意见
"""
import io
from datetime import datetime
from typing import Any


def render_opinions_docx(
    project_name: str,
    country: str,
    amount: str,
    stage: str,
    opinions: list[dict[str, Any]],
    review_time: str | None = None,
) -> bytes:
    """生成评审意见 Word 文档

    Args:
        project_name: 项目名称
        country: 国别
        amount: 合同金额
        stage: 阶段
        opinions: 评审意见列表 [{department, rule_id, title, content, risk_level, evidence}]
        review_time: 评审时间

    Returns:
        bytes: .docx 文件内容
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 标题
    title = doc.add_heading('', level=0)
    run = title.add_run('海外项目评审意见')
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 项目信息
    doc.add_paragraph('')
    info = doc.add_table(rows=5, cols=2)
    info.style = 'Light Grid Accent 1'
    now = review_time or datetime.now().strftime('%Y-%m-%d %H:%M')
    for i, (k, v) in enumerate([('项目名称', project_name), ('国别', country), ('合同金额', amount), ('阶段', stage), ('评审时间', now)]):
        info.rows[i].cells[0].text = k
        info.rows[i].cells[1].text = v

    doc.add_paragraph('')

    # 风险等级标注
    level_colors = {'high': RGBColor(0xCC, 0x00, 0x00), 'medium': RGBColor(0xCC, 0x88, 0x00), 'low': RGBColor(0x00, 0x66, 0x00)}

    # 按部门分组
    from collections import defaultdict
    groups: dict[str, list] = defaultdict(list)
    for op in opinions:
        dept = op.get('department', '综合') or '综合'
        groups[dept].append(op)

    dept_order = sorted(groups.keys(), key=lambda d: len(groups[d]), reverse=True)

    for dept in dept_order:
        items = groups[dept]
        doc.add_heading(dept, level=2)

        for idx, op in enumerate(items, 1):
            risk = op.get('risk_level', 'medium')
            color = level_colors.get(risk, level_colors['medium'])

            # 规则标题
            p = doc.add_paragraph()
            run = p.add_run(f'{idx}. {op.get("rule_id", "")}: {op.get("title", "")[:80]}')
            run.bold = True
            run.font.size = Pt(11)

            # 风险等级标签
            run2 = p.add_run(f'  [{risk.upper()}]')
            run2.font.color.rgb = color
            run2.font.size = Pt(9)

            # 证据内容
            evidence = op.get('evidence', '') or op.get('content', '')
            if evidence:
                ep = doc.add_paragraph(f'判定依据: {evidence[:500]}')
                ep.paragraph_format.space_before = Pt(2)
                ep.paragraph_format.space_after = Pt(6)
                for run in ep.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        doc.add_paragraph('')

    # 页脚
    doc.add_paragraph('— 本报告由 AI 评审系统自动生成，仅供参考 —').alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
